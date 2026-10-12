from src.lab05.json_csv import check_extension, read_csv


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Переносит CSV в Sheet1; ширина колонок не меньше 8."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    check_extension(xlsx_path, ".xlsx")
    header, rows = read_csv(csv_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    for i in range(len(header)):
        width = 8
        for row in [header] + rows:
            width = max(width, len(str(row[i])) + 2)
        sheet.column_dimensions[get_column_letter(i + 1)].width = width
    workbook.save(xlsx_path)
